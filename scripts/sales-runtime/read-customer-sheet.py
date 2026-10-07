"""Read a pinned workbook, including hidden rows. Never save or modify it.

stdout contains sensitive source data, intended ONLY for the review runner's
captured pipe. Run run-sheet-import-review.mjs, not this helper directly.
"""
import argparse
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import warnings

from openpyxl import load_workbook


def extract(path, expected_sha256, snapshot_date):
    if not re.fullmatch(r"[a-f0-9]{64}", expected_sha256):
        raise ValueError("SOURCE_HASH_REQUIRED")
    if date.fromisoformat(snapshot_date).isoformat() != snapshot_date:
        raise ValueError("SNAPSHOT_DATE_INVALID")
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError("SOURCE_TOO_LARGE")
    data = path.read_bytes()
    actual_hash = hashlib.sha256(data).hexdigest()
    if actual_hash != expected_sha256:
        raise ValueError("SOURCE_HASH_CHANGED")
    # Known Google export has an invalid pivot cache dependency, outside A:Q.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*invalid dependency definitions.*")
        values_book = load_workbook(BytesIO(data), data_only=True)
        formulas_book = load_workbook(BytesIO(data), data_only=False)
    sheet = "ข้อมูลลูกค้า"
    values, formulas = values_book[sheet], formulas_book[sheet]
    if values.max_row > 50000 or values.max_row < 2:
        raise ValueError("SOURCE_ROW_LIMIT")

    def primitive(value):
        if isinstance(value, datetime):
            # Preserve any non-midnight time for explicit review, do not truncate.
            return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        raise ValueError("UNSUPPORTED_CELL_TYPE")

    result = {
        "version": "customer-sheet-extract-v1", "sheet": sheet,
        "sha256": actual_hash, "snapshotDate": snapshot_date,
        "phonePolicy": "unknown_confirmed",
        "headers": [primitive(values.cell(1, col).value) for col in range(1, 18)],
        "maxRow": values.max_row,
        "rows": [{"rowNumber": row, "hidden": bool(formulas.row_dimensions[row].hidden),
                  "values": [primitive(values.cell(row, col).value) for col in range(1, 18)],
                  "formulaColumns": [col - 1 for col in range(1, 18)
                                     if formulas.cell(row, col).data_type == "f"]}
                 for row in range(2, values.max_row + 1)],
    }
    values_book.close()
    formulas_book.close()
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True)
    parser.add_argument("--sha256", required=True)
    parser.add_argument("--snapshot-date", required=True)
    args = parser.parse_args()
    try:
        result = extract(Path(args.source), args.sha256, args.snapshot_date)
        sys.stdout.reconfigure(encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False, allow_nan=False))
    except Exception:
        # No file paths, cell content, customer names or traceback in diagnostics.
        print("CUSTOMER_SHEET_EXTRACTION_FAILED", file=sys.stderr)
        sys.exit(1)
